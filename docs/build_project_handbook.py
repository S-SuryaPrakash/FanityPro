from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Content_Filter_Project_Handbook.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B76"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "C28C18"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "111111"
TABLE_WIDTH = 9360
TABLE_INDENT = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=None, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margins=CELL_MARGINS):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent=TABLE_INDENT):
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, value in (("tblW", TABLE_WIDTH), ("tblInd", indent)):
        element = tbl_pr.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tbl_pr.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_numbering_definition(document, bullet: bool) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_styles(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = document.styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_after = Pt(18)

    code_style = document.styles.add_style("Code Block", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = rgb(NAVY)
    code_style.paragraph_format.left_indent = Inches(0.15)
    code_style.paragraph_format.right_indent = Inches(0.15)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.1

    path_style = document.styles.add_style("File Path", 1)
    path_style.font.name = "Consolas"
    path_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    path_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    path_style.font.size = Pt(8.5)
    path_style.font.color.rgb = rgb(MUTED)
    path_style.paragraph_format.space_after = Pt(4)

    for header in (section.header,):
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("CONTENT FILTER PROJECT HANDBOOK  |  LEARNING EDITION")
        set_run_font(r, size=8.5, color=MUTED, bold=True)
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Content Filter  |  Page ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(footer)
    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = first_footer.add_run("Source snapshot: Module 7 complete | Human evaluation gate remains open")
    set_run_font(r, size=8.5, color=MUTED)


def shade_paragraph(paragraph, fill=CALLOUT, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        borders.append(left)
        p_pr.append(borders)


def add_callout(document, label: str, text: str, kind="info"):
    colors = {"info": (CALLOUT, BLUE), "warning": ("FFF8E8", GOLD), "risk": ("FDECEC", RED)}
    fill, border = colors[kind]
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, fill, border)
    r = p.add_run(label + ": ")
    set_run_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    return p


def add_code(document, text: str):
    p = document.add_paragraph(style="Code Block")
    shade_paragraph(p, LIGHT_GRAY)
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=9, color=NAVY)
    return p


def add_bullet(document, text: str, bullet_num_id: int):
    p = document.add_paragraph()
    apply_numbering(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    return p


def add_number(document, text: str, decimal_num_id: int):
    p = document.add_paragraph()
    apply_numbering(p, decimal_num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    return p


def add_labeled_paragraph(document, label: str, text: str):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + ": ")
    set_run_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    return p


def add_table(document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    mark_repeat_header(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, size=9.5, color=BLACK)
    set_table_geometry(table, widths)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_class_entry(document, name: str, path: str, purpose: str, inputs: str, outputs: str,
                    works_with: str, logic: list[str], takeaway: str, bullet_num_id: int,
                    status: str | None = None):
    document.add_heading(name, level=3)
    p = document.add_paragraph(style="File Path")
    p.add_run(path)
    if status:
        add_callout(document, "Current status", status, "warning" if "legacy" in status.lower() or "unused" in status.lower() else "info")
    add_labeled_paragraph(document, "Purpose", purpose)
    add_labeled_paragraph(document, "Receives", inputs)
    add_labeled_paragraph(document, "Produces or changes", outputs)
    add_labeled_paragraph(document, "Works with", works_with)
    if logic:
        p = document.add_paragraph()
        r = p.add_run("Important logic")
        set_run_font(r, bold=True, color=NAVY)
        p.paragraph_format.space_after = Pt(2)
        for item in logic:
            add_bullet(document, item, bullet_num_id)
    add_callout(document, "Beginner takeaway", takeaway, "info")


JAVA_CLASSES = {
    "Application entry and configuration": [
        dict(name="ContentfilterApplication", path="src/main/java/com/example/contentfilter/ContentfilterApplication.java",
             purpose="Starts Spring Boot and tells Spring to scan the project for components and configuration-property records.",
             inputs="Command-line arguments and the application configuration available at startup.",
             outputs="A running embedded web application with controllers, services, filters, and configuration beans wired together.",
             works_with="SpringApplication, @SpringBootApplication, and @ConfigurationPropertiesScan.",
             logic=["main() delegates startup to SpringApplication.run().", "Component scanning discovers @Service, @Component, @Configuration, @RestController, and @RestControllerAdvice classes."],
             takeaway="This is the front door of the Java process. It starts the container; it does not classify messages itself."),
        dict(name="UploadLimitsProperties", path="src/main/java/com/example/contentfilter/config/UploadLimitsProperties.java",
             purpose="Holds validated resource limits for synchronous workbook processing.",
             inputs="content-filter.upload.* values from application.properties or environment variables.",
             outputs="An immutable settings record injected into ExcelService.",
             works_with="ExcelService and Spring configuration-property binding.",
             logic=["Validates positive file size and processing time.", "Carries worksheet, physical-row, sequence, cell, per-sequence text, total-text, and time limits."],
             takeaway="Limits belong in configuration so deployments can be made safer without editing business code."),
        dict(name="RiskPolicyProperties", path="src/main/java/com/example/contentfilter/config/RiskPolicyProperties.java",
             purpose="Stores the versioned decision thresholds, uncertainty margin, and category precedence.",
             inputs="content-filter.classification.policy.* configuration.",
             outputs="A validated policy configuration consumed by VersionedRiskPolicy.",
             works_with="RiskCategory and VersionedRiskPolicy.",
             logic=["Requires a threshold for every detected-risk category.", "Requires each precedence category exactly once.", "Rejects missing, duplicate, non-finite, or out-of-range values."],
             takeaway="The model supplies scores; this configuration defines how the business interprets them."),
        dict(name="ModelServiceProperties", path="src/main/java/com/example/contentfilter/config/ModelServiceProperties.java",
             purpose="Defines how Java connects to FastAPI and which exact model identity it will trust.",
             inputs="Model-service URL, expected model ID and 40-character revision, timeouts, batch size, attempts, and provisional-model switch.",
             outputs="Validated settings for FastApiClientConfiguration and FastApiRiskModel.",
             works_with="FastApiRiskModel and FastApiClientConfiguration.",
             logic=["Accepts only an HTTP(S) origin without credentials, query, fragment, or custom path.", "Limits connect/read timeouts and retry attempts.", "Ensures all attempts fit inside the 60-second V1 request budget.", "Pins the model revision so a remote deployment cannot silently change models."],
             takeaway="This class is a safety contract, not merely a list of URLs."),
        dict(name="FastApiClientConfiguration", path="src/main/java/com/example/contentfilter/config/FastApiClientConfiguration.java",
             purpose="Creates the private RestClient used to call FastAPI.",
             inputs="ModelServiceProperties.",
             outputs="A qualified fastApiRestClient bean.",
             works_with="JDK HttpClient, Spring RestClient, and FastApiRiskModel.",
             logic=["Created only when the provider property is fastapi.", "Disables redirects, uses HTTP/1.1, and applies connect/read timeouts.", "Sets the FastAPI base URL and JSON Accept header."],
             takeaway="Centralising HTTP-client construction keeps timeout and security behavior consistent."),
    ],
    "HTTP controllers": [
        dict(name="FileClassificationController", path="src/main/java/com/example/contentfilter/controller/FileClassificationController.java",
             purpose="Exposes the complete versioned upload-to-annotated-workbook workflow.",
             inputs="A multipart form-data request whose file part contains an .xlsx workbook.",
             outputs="A binary .xlsx attachment named classified-<original>.xlsx.",
             works_with="WorkbookClassificationService and ClassifiedWorkbook.",
             logic=["POST /api/v1/files/classify is the production-shaped V1 endpoint.", "Marks the download no-store/no-cache because conversation content may be sensitive.", "Adds a UTF-8 attachment filename, content length, Excel content type, and nosniff header.", "Keeps HTTP adaptation in the controller and business logic in services."],
             takeaway="The controller is intentionally thin: receive HTTP, call one use-case service, build HTTP response."),
        dict(name="UploadController", path="src/main/java/com/example/contentfilter/controller/UploadController.java",
             purpose="Exposes the original JSON upload prototype.",
             inputs="A multipart file at POST /upload.",
             outputs="UploadResponse JSON containing one legacy classification per non-empty row.",
             works_with="UploadClassificationService.",
             logic=["Returns file metadata and row results.", "Delegates validation, Excel parsing, and classification rather than doing them in the controller."],
             takeaway="This endpoint is retained for compatibility while the versioned workbook endpoint replaces it.",
             status="Legacy compatibility path. New product behavior should use /api/v1/files/classify."),
    ],
    "Domain records and enums": [
        dict(name="ExtractedSequence", path="src/main/java/com/example/contentfilter/domain/ExtractedSequence.java",
             purpose="Represents one complete message together with the exact Excel cells from which it came.",
             inputs="Sequence ID, zero-based sheet/row/column coordinates, text, and optional context.",
             outputs="An immutable, validated domain object used from extraction through reporting.",
             works_with="ExcelService, RiskModel, ConversationContext, and ExcelReportService.",
             logic=["Rejects blank IDs/text and negative or empty coordinates.", "Defensively copies source columns.", "Uses an empty context when none is supplied."],
             takeaway="This object is the bridge between Excel and classification; coordinates let the report find the original cell later."),
        dict(name="ConversationContext", path="src/main/java/com/example/contentfilter/domain/ConversationContext.java",
             purpose="Carries optional conversation metadata without adding it to the text sent to the model.",
             inputs="Conversation ID, message ID, speaker role, timestamp, language, and channel.",
             outputs="Normalised immutable context; blanks become null.",
             works_with="ExtractedSequence and FastApiRiskModel.",
             logic=["Provides a reusable empty() instance.", "Strips surrounding whitespace from present values."],
             takeaway="Context stays attached for audit/reporting while model input remains the intended message text only."),
        dict(name="ConversationColumnMapping", path="src/main/java/com/example/contentfilter/domain/ConversationColumnMapping.java",
             purpose="Describes which Excel headers represent message text and optional metadata.",
             inputs="Header-row index and header names.",
             outputs="A validated mapping used by ExcelService.",
             works_with="ExcelService and WorkbookClassificationService.",
             logic=["Header matching is case-insensitive and trims spaces.", "Requires a text header and prevents one column from being assigned to multiple fields.", "standard() defines row 0 with text plus the documented optional headers."],
             takeaway="The mapping separates the workbook's physical column names from Java business concepts."),
        dict(name="ModelPrediction", path="src/main/java/com/example/contentfilter/domain/ModelPrediction.java",
             purpose="Stores provider-neutral model evidence before any business decision is applied.",
             inputs="Sequence ID, risk-score map, truncation flag, model ID, and model revision.",
             outputs="Validated immutable score evidence.",
             works_with="RiskModel, RiskClassificationService, and VersionedRiskPolicy.",
             logic=["Allows only the five detected-risk categories.", "Requires each supplied score to be finite and between 0 and 1.", "Preserves model identity for auditability."],
             takeaway="Prediction is evidence. It is deliberately different from ClassificationResult, which is a policy decision."),
        dict(name="ClassificationResult", path="src/main/java/com/example/contentfilter/domain/ClassificationResult.java",
             purpose="Represents the final auditable Java policy decision for one message.",
             inputs="Model evidence interpreted by VersionedRiskPolicy.",
             outputs="Primary category, severity, confidence, all scores, review status/reason, model identity, and policy version.",
             works_with="VersionedRiskPolicy, ConversationRiskAssessment, and ExcelReportService.",
             logic=["Requires valid confidence and model/policy versions.", "Requires a human-readable reason whenever review is required.", "Defensively copies score evidence."],
             takeaway="This is the answer the business uses, while ModelPrediction is what the model observed."),
        dict(name="ConversationRiskAssessment", path="src/main/java/com/example/contentfilter/domain/ConversationRiskAssessment.java",
             purpose="Keeps extracted source messages and classification decisions together in matching order.",
             inputs="A sequence list and a result list.",
             outputs="An immutable correlated assessment for report generation.",
             works_with="ConversationRiskWorkflowService and ExcelReportService.",
             logic=["Requires equal list sizes.", "Checks every sequence ID against the result at the same position."],
             takeaway="This record prevents a classification from being written onto the wrong Excel row."),
        dict(name="ClassifiedWorkbook", path="src/main/java/com/example/contentfilter/domain/ClassifiedWorkbook.java",
             purpose="Packages the final download filename and workbook bytes.",
             inputs="A valid .xlsx filename and non-empty byte array.",
             outputs="An immutable download object.",
             works_with="WorkbookClassificationService and FileClassificationController.",
             logic=["Clones the byte array on construction and access so callers cannot mutate it.", "Requires a non-empty .xlsx result."],
             takeaway="Even byte arrays are mutable; defensive copies protect the completed report."),
        dict(name="RiskCategory", path="src/main/java/com/example/contentfilter/domain/RiskCategory.java",
             purpose="Defines the stable V1 risk vocabulary shared across Java, Python, policy, tests, and reports.",
             inputs="No runtime input; it is an enum.",
             outputs="Five detected-risk values plus NO_AUTOMATED_FLAG and MANUAL_REVIEW outcomes.",
             works_with="Every model, policy, configuration, and report component.",
             logic=["isDetectedRisk() excludes policy outcomes from model-score maps."],
             takeaway="A stable taxonomy prevents every model provider from inventing a different application contract."),
        dict(name="RiskSeverity", path="src/main/java/com/example/contentfilter/domain/RiskSeverity.java",
             purpose="Defines human-review priority levels.",
             inputs="Assigned by VersionedRiskPolicy.",
             outputs="NONE, LOW, MEDIUM, HIGH, or CRITICAL.",
             works_with="ClassificationResult, VersionedRiskPolicy, and review-queue ordering.",
             logic=["Severity is a policy result, not a direct neural-network output."],
             takeaway="Category says what kind of concern exists; severity says how urgently people should review it."),
    ],
    "Legacy DTOs": [
        dict(name="ClassificationRequest", path="src/main/java/com/example/contentfilter/dto/ClassificationRequest.java",
             purpose="A minimal text request record from the early prototype.", inputs="A text string.", outputs="A transport object only.",
             works_with="No active controller in the current workflow.", logic=[],
             takeaway="Not every file is active. This DTO can be removed or reused deliberately in a later cleanup.",
             status="Currently unused by the active endpoints."),
        dict(name="ClassificationResponse", path="src/main/java/com/example/contentfilter/dto/ClassificationResponse.java",
             purpose="Carries category, random demonstration confidence, and timestamp for the legacy classifier.",
             inputs="KeywordClassifier output.", outputs="Legacy JSON classification data.",
             works_with="KeywordClassifier and RowClassificationResponse.",
             logic=["categoryLower() returns a copy with a lower-case category."],
             takeaway="This DTO belongs to the old three-category prototype, not the production risk taxonomy.",
             status="Legacy compatibility DTO."),
        dict(name="RowClassificationResponse", path="src/main/java/com/example/contentfilter/dto/RowClassificationResponse.java",
             purpose="Adds the one-based Excel row number and source text to a legacy ClassificationResponse.",
             inputs="An ExtractedSequence plus legacy classification.", outputs="One item in UploadResponse.results.",
             works_with="UploadClassificationService and UploadResponse.", logic=[],
             takeaway="The ClassificationResponse field is supplied by KeywordClassifier through the legacy service interface.",
             status="Legacy compatibility DTO."),
        dict(name="UploadRequest", path="src/main/java/com/example/contentfilter/dto/UploadRequest.java",
             purpose="An early placeholder request record containing optional text.", inputs="Optional text.", outputs="A transport object only.",
             works_with="No active controller.", logic=[],
             takeaway="It is safe to treat this as unused scaffolding, not part of the current program flow.",
             status="Currently unused by the active endpoints."),
        dict(name="UploadResponse", path="src/main/java/com/example/contentfilter/dto/UploadResponse.java",
             purpose="Returns file metadata and legacy row classifications as JSON.",
             inputs="File name, size, MIME type, and row results.", outputs="The body of POST /upload.",
             works_with="UploadController and RowClassificationResponse.", logic=[],
             takeaway="The versioned endpoint returns Excel bytes instead, so this DTO remains only for compatibility.",
             status="Legacy compatibility DTO."),
    ],
    "Interfaces (ports)": [
        dict(name="RiskModel", path="src/main/java/com/example/contentfilter/service/RiskModel.java",
             purpose="Defines the model-provider boundary.", inputs="A batch of ExtractedSequence objects.", outputs="A list of ModelPrediction evidence.",
             works_with="DeterministicRiskModel, FastApiRiskModel, and RiskClassificationService.",
             logic=["Java orchestration depends on this interface rather than a specific provider."],
             takeaway="An interface is a plug shape: both the deterministic adapter and FastAPI adapter fit the same socket."),
        dict(name="ClassificationPolicy", path="src/main/java/com/example/contentfilter/service/ClassificationPolicy.java",
             purpose="Defines how model evidence becomes an auditable decision.", inputs="One ModelPrediction.", outputs="One ClassificationResult.",
             works_with="VersionedRiskPolicy and RiskClassificationService.", logic=[],
             takeaway="Keeping policy behind an interface makes model evidence and business rules separately testable."),
        dict(name="ClassificationService", path="src/main/java/com/example/contentfilter/service/ClassificationService.java",
             purpose="Defines the batch-first application classification use case.", inputs="Extracted sequences.", outputs="Correlated classification results.",
             works_with="RiskClassificationService and ConversationRiskWorkflowService.",
             logic=["The contract says callers must not trust provider response order."],
             takeaway="This is the production-shaped classifier interface used by the new workflow."),
        dict(name="LegacyClassificationService", path="src/main/java/com/example/contentfilter/service/LegacyClassificationService.java",
             purpose="Keeps the old one-string keyword classifier replaceable while migration continues.", inputs="One text string.", outputs="ClassificationResponse.",
             works_with="KeywordClassifier and UploadClassificationService.", logic=[],
             takeaway="The name warns new code not to build on the obsolete single-text contract.",
             status="Legacy interface; new code uses ClassificationService."),
    ],
    "Application services and adapters": [
        dict(name="ExcelService", path="src/main/java/com/example/contentfilter/service/ExcelService.java",
             purpose="Safely validates .xlsx uploads and extracts bounded, coordinate-aware messages from the first worksheet.",
             inputs="MultipartFile, optionally with ConversationColumnMapping.", outputs="An immutable list of ExtractedSequence objects.",
             works_with="Apache POI, UploadLimitsProperties, WorkbookProcessingException, and ConversationContext.",
             logic=["Checks file presence, size, filename, extension, MIME type, and real OOXML signature.", "Rejects encrypted, macro-enabled, malformed, over-limit, or unsupported workbooks.", "Configures zip-bomb defenses through ZipSecureFile.", "Uses cached formula display values rather than recalculating or following external links.", "In mapped mode, skips the header row, sends only the text column to classification, and retains optional context.", "Records physical zero-based coordinates and enforces row/cell/text/time limits."],
             takeaway="Apache POI belongs here because Excel parsing is business infrastructure, not HTTP-controller logic."),
        dict(name="KeywordClassifier", path="src/main/java/com/example/contentfilter/service/KeywordClassifier.java",
             purpose="Implements the original demonstration classifier using simple substring rules.",
             inputs="One non-blank string up to 10,000 characters.", outputs="abusive, professional, or neutral with random placeholder confidence and timestamp.",
             works_with="LegacyClassificationService and InvalidClassificationRequestException.",
             logic=["stupid/idiot -> abusive; regards -> professional; otherwise neutral.", "Matching is case-insensitive.", "Random confidence is not a real probability."],
             takeaway="This class is useful for learning and compatibility but must not be mistaken for the Hugging Face model.",
             status="Legacy demonstration implementation."),
        dict(name="UploadClassificationService", path="src/main/java/com/example/contentfilter/service/UploadClassificationService.java",
             purpose="Coordinates the old row-as-message JSON workflow.", inputs="MultipartFile.", outputs="Immutable row-level legacy results.",
             works_with="ExcelService and LegacyClassificationService.",
             logic=["Extracts simple rows, loops through them, classifies each string individually, and converts zero-based row coordinates to one-based user-facing row numbers."],
             takeaway="This is a coordinating service: it joins two capabilities but does not parse Excel or implement classification itself.",
             status="Legacy compatibility workflow."),
        dict(name="DeterministicRiskModel", path="src/main/java/com/example/contentfilter/service/DeterministicRiskModel.java",
             purpose="Provides predictable multi-label scores so the production-shaped Java workflow can be developed without a running model server.",
             inputs="A sequence batch.", outputs="ModelPrediction objects with fixed scores based on known phrases.",
             works_with="RiskModel and RiskClassificationService.",
             logic=["Selected by default through content-filter.classification.provider=deterministic.", "Starts every detected-risk score at 0.01 and raises matching categories to fixed values.", "Always reports a deterministic development model ID/revision and no truncation."],
             takeaway="A test adapter lets you exercise the real architecture without pretending it is a trained model."),
        dict(name="FastApiRiskModel", path="src/main/java/com/example/contentfilter/service/FastApiRiskModel.java",
             purpose="Implements RiskModel by calling the private FastAPI batch endpoint safely.",
             inputs="Extracted sequences, a qualified RestClient, and ModelServiceProperties.", outputs="Validated provider-neutral ModelPrediction objects.",
             works_with="FastAPI /ready and /api/v1/classify/batch, CorrelationIdFilter, and ModelServiceException.",
             logic=["Splits work into configured batches and restores the original request order by sequence ID.", "Sends conversation ID, speaker role, and language as context without changing text.", "Propagates correlation IDs from the logging context.", "Retries at most one transient network/502/503/504 failure and never silently falls back to another model.", "Validates exact model ID, pinned revision, release-gate metadata, prediction IDs, score categories, and truncation flags.", "Rejects provisional evidence unless the explicit development waiver is enabled."],
             takeaway="This adapter assumes the remote service can fail or lie accidentally, so it validates every important part of the response."),
        dict(name="RiskClassificationService", path="src/main/java/com/example/contentfilter/service/RiskClassificationService.java",
             purpose="Combines model evidence with policy while protecting request/result correlation.",
             inputs="A batch of ExtractedSequence objects.", outputs="One ClassificationResult for every requested sequence in original order.",
             works_with="RiskModel and ClassificationPolicy.",
             logic=["Rejects null or duplicate sequence IDs.", "Indexes provider predictions by ID rather than trusting their order.", "Fails if IDs are missing, extra, or duplicated.", "Applies the policy to the correct prediction for each original sequence."],
             takeaway="Correlation checks prevent the most dangerous reporting bug: putting one person's result on another row."),
        dict(name="VersionedRiskPolicy", path="src/main/java/com/example/contentfilter/service/VersionedRiskPolicy.java",
             purpose="Turns multi-label scores into the V1 category, severity, and human-review decision.",
             inputs="ModelPrediction plus RiskPolicyProperties.", outputs="ClassificationResult.",
             works_with="ClassificationPolicy, RiskPolicyProperties, RiskCategory, and RiskSeverity.",
             logic=["Truncated input always goes to high-priority manual review.", "Two passing scores within the uncertainty margin become MANUAL_REVIEW.", "Otherwise the first passing category in configured precedence becomes primary.", "A score just below a threshold becomes low-priority manual review.", "When no threshold is approached, the result is NO_AUTOMATED_FLAG, not a safety guarantee.", "Threat is critical; identity attack/harassment high; profanity/toxicity medium."],
             takeaway="Thresholds and human-review policy belong in deterministic Java code so model replacement does not silently change operational decisions."),
        dict(name="ConversationRiskWorkflowService", path="src/main/java/com/example/contentfilter/service/ConversationRiskWorkflowService.java",
             purpose="Coordinates mapped extraction and batch risk classification.", inputs="MultipartFile and ConversationColumnMapping.", outputs="ConversationRiskAssessment.",
             works_with="ExcelService and ClassificationService.",
             logic=["Extracts messages with coordinates/context, classifies the batch, then bundles both lists into a correlation-checking assessment."],
             takeaway="This service is the center of the classification workflow, but report generation remains a separate concern."),
        dict(name="ExcelReportService", path="src/main/java/com/example/contentfilter/service/ExcelReportService.java",
             purpose="Reopens the original workbook and creates the auditable classified copy.",
             inputs="Original MultipartFile, column mapping, and ConversationRiskAssessment.", outputs="Final workbook bytes.",
             works_with="Apache POI, ExtractedSequence coordinates, ClassificationResult, and ReportGenerationException.",
             logic=["Rejects pre-existing reserved sheets or Content Filter columns rather than overwriting user data.", "Creates reusable styles by primary category instead of one style per cell.", "Colours the original source message cells.", "Appends category, severity, confidence, review, model, policy, and five score columns.", "Creates Review Queue ordered by severity then confidence.", "Creates Summary counts and a Legend with limitations.", "Writes to memory and returns a byte array; it does not permanently store conversation content."],
             takeaway="The report service uses the coordinates saved much earlier by ExcelService; that is why the domain model preserved them."),
        dict(name="WorkbookClassificationService", path="src/main/java/com/example/contentfilter/service/WorkbookClassificationService.java",
             purpose="Runs the complete versioned use case from upload to downloadable workbook.",
             inputs="MultipartFile.", outputs="ClassifiedWorkbook.",
             works_with="ConversationRiskWorkflowService, ExcelReportService, and ConversationColumnMapping.standard().",
             logic=["Rejects a workbook with no non-empty messages.", "Builds a safe classified-<name>.xlsx filename, normalises the extension, removes control characters, and limits length.", "Keeps HTTP concerns out of the workflow."],
             takeaway="This is the one method the versioned controller calls because it represents one complete business use case."),
        dict(name="FastApiModelHealthIndicator", path="src/main/java/com/example/contentfilter/service/FastApiModelHealthIndicator.java",
             purpose="Makes Spring readiness depend on FastAPI when that provider is selected.",
             inputs="FastApiRiskModel.readiness().", outputs="Actuator UP, OUT_OF_SERVICE, or DOWN health.",
             works_with="Spring Boot Actuator and FastApiRiskModel.",
             logic=["Reports safe model/revision/release metadata when ready.", "Maps known model-service failures to OUT_OF_SERVICE and unexpected Java failures to DOWN."],
             takeaway="Liveness means Java is running; readiness means Java can safely accept classification work."),
    ],
    "Exceptions and web infrastructure": [
        dict(name="WorkbookProcessingException", path="src/main/java/com/example/contentfilter/exception/WorkbookProcessingException.java",
             purpose="Carries safe workbook failure categories from services to the HTTP error translator.",
             inputs="Reason, client-facing message, and optional cause.", outputs="A typed runtime exception.",
             works_with="ExcelService, WorkbookClassificationService, ExcelReportService, and GlobalExceptionHandler.",
             logic=["Reasons distinguish invalid request, unsupported type, limit exceeded, and invalid workbook."],
             takeaway="A reason enum is more reliable than parsing exception-message text to decide an HTTP status."),
        dict(name="InvalidClassificationRequestException", path="src/main/java/com/example/contentfilter/exception/InvalidClassificationRequestException.java",
             purpose="Signals blank or oversized input to the legacy KeywordClassifier.", inputs="A safe validation message.", outputs="A runtime exception mapped to HTTP 400.",
             works_with="KeywordClassifier and GlobalExceptionHandler.", logic=[],
             takeaway="This exception belongs to the legacy classifier path."),
        dict(name="ReportGenerationException", path="src/main/java/com/example/contentfilter/exception/ReportGenerationException.java",
             purpose="Signals an internal failure while producing the annotated workbook without leaking Apache POI details.",
             inputs="Safe message and optional cause.", outputs="A runtime exception mapped to REPORT_GENERATION_FAILED.",
             works_with="ExcelReportService and GlobalExceptionHandler.", logic=[],
             takeaway="Internal details are logged with a correlation ID; callers receive a stable safe error."),
        dict(name="ModelServiceException", path="src/main/java/com/example/contentfilter/service/ModelServiceException.java",
             purpose="Classifies safe failures from the private FastAPI contract.",
             inputs="Reason, safe message, and optional cause.", outputs="A typed model-service runtime exception.",
             works_with="FastApiRiskModel, FastApiModelHealthIndicator, and GlobalExceptionHandler.",
             logic=["Reasons cover unavailable, timeout, request rejected, invalid response, and model not approved."],
             takeaway="Callers need different HTTP behavior for a timeout, a bad remote contract, and an unapproved model."),
        dict(name="GlobalExceptionHandler", path="src/main/java/com/example/contentfilter/exception/GlobalExceptionHandler.java",
             purpose="Converts Java exceptions into consistent RFC 9457-style Problem Details responses.",
             inputs="Exceptions and HttpServletRequest.", outputs="HTTP status plus safe problem JSON containing errorCode, correlationId, and timestamp.",
             works_with="Every controller, domain exception, CorrelationIdFilter, and Spring @RestControllerAdvice.",
             logic=["Maps workbook errors to 400/413/415, model timeouts to 504, model outages/unapproved state to 503, invalid model evidence to 502, report failures to 500, and unexpected errors to a generic 500.", "Never exposes a remote FastAPI body or internal stack trace to the caller.", "Logs unexpected/report failures with the correlation ID."],
             takeaway="Controllers stay clean because all error-to-HTTP translation happens in one place."),
        dict(name="CorrelationIdFilter", path="src/main/java/com/example/contentfilter/web/CorrelationIdFilter.java",
             purpose="Assigns a traceable identifier to every Java HTTP request.",
             inputs="Optional X-Correlation-ID header.", outputs="Safe response header, request attribute, and MDC logging value.",
             works_with="GlobalExceptionHandler and FastApiRiskModel.",
             logic=["Accepts only 1-64 safe letters/digits/dot/underscore/hyphen.", "Generates a UUID when absent or unsafe.", "Uses OncePerRequestFilter and highest precedence.", "Automatically clears the MDC value after the request."],
             takeaway="When a user reports an error, the correlation ID joins the Java log, FastAPI call, and HTTP response."),
    ],
}


PYTHON_MODULES = [
    ("app/__init__.py", "Marks app as the internal model-service Python package. It contains only a package-level description."),
    ("app/settings.py", "ModelServiceSettings reads CONTENT_FILTER_MODEL_* environment variables, fixes the candidate key/manifest/cache path and resource limits, and deliberately restricts release state to PROVISIONAL with approved_for_production=False."),
    ("app/schemas.py", "Defines strict Pydantic request/response models: five RiskCategory values, bounded English SequenceInput, unique BatchClassificationRequest, exact RiskScores, prediction/batch/health responses, validation issues, and Problem Details. Extra JSON fields are rejected."),
    ("app/model_manifest.py", "Loads one eligible model from candidate-models.json, requires a 40-character commit revision, verifies exact domain-label coverage, normalises native labels, and fails closed when model metadata is incomplete."),
    ("app/runtime.py", "Defines the replaceable RiskModelRuntime protocol and the real HuggingFaceRiskRuntime. It loads a pinned tokenizer/model once, forbids remote code, requires safetensors, validates native labels, caps token length at 512, uses torch.inference_mode(), applies sigmoid, maps native labels to five domain scores, records truncation, and serialises CPU inference with a lock."),
    ("app/service.py", "ClassificationApplicationService enforces runtime batch/text limits, checks model readiness, validates that every sequence ID returns exactly once, checks all five finite 0-1 scores, and restores request order. It never applies Java business policy."),
    ("app/main.py", "Creates FastAPI with injectable runtime factories. Lifespan loads the model once; middleware manages correlation IDs; exception handlers return safe problems; /live checks the process, /ready checks the loaded model, and /api/v1/classify/batch returns model evidence and release metadata."),
]


EVALUATION_SCRIPTS = [
    ("scripts/build_domain_dataset.py", "Deterministically generates the 240-row privacy-safe synthetic calibration/test corpus. It includes benign, identity, quoted-harm, context-sensitive, threat, hate, harassment, profanity, toxicity, multi-label, human-like, and AI-like slices. Every generated row remains draft."),
    ("scripts/evaluation_lib.py", "Shared dependency-free functions for JSONL loading, schema validation, duplicate detection, coverage gates, release-readiness checks, and per-label/micro/macro/slice metrics including precision, recall, F1, FPR, FNR, and latency."),
    ("scripts/validate_dataset.py", "Command-line validator with basic, domain, and release profiles. The release profile intentionally fails until every example is genuinely adjudicated."),
    ("scripts/run_hf_evaluation.py", "Loads one eligible pinned Hugging Face candidate, runs a chosen calibration/test split in batches, maps native scores to application categories, and writes prediction JSONL with model identity, truncation, and latency."),
    ("scripts/select_thresholds.py", "Uses calibration data only to select a threshold per label. It maximises F2 subject to minimum precision when feasible and writes a draft, never-production-approved threshold document."),
    ("scripts/evaluate_predictions.py", "Validates exact prediction IDs and thresholds, evaluates the untouched split, and writes an auditable report containing model/dataset/threshold versions and multilabel metrics."),
]


JAVA_TESTS = [
    ("ContentfilterApplicationTests", "Spring context starts successfully."),
    ("UploadLimitsPropertiesTests", "Invalid non-positive file/time limits are rejected."),
    ("ModelServicePropertiesTests", "Unsafe model URLs and request budgets are rejected."),
    ("OperationalContractTests", "Actuator probes, correlation-aware Problem Details, and configured upload limits work."),
    ("UploadControllerTests", "Legacy multipart name, workbook errors, and physical row-number output."),
    ("FileClassificationControllerTests", "The versioned endpoint returns a non-cacheable Excel attachment and rejects missing/empty uploads."),
    ("ExcelServiceTests", "Coordinates, blank rows, cached formulas, mappings, signatures, macros, malformed OOXML, resource limits, and zip-bomb safeguards."),
    ("KeywordClassifierTests", "Legacy categories, metadata, and blank-text rejection."),
    ("RiskClassificationServiceTests", "Out-of-order predictions are correlated by ID and mismatches fail."),
    ("VersionedRiskPolicyTests", "Threat, no-flag, near-threshold, conflict, and truncation decisions."),
    ("ConversationRiskWorkflowServiceTests", "Mapped workbook extraction reaches the decision policy with context and correlation intact."),
    ("FastApiRiskModelTests", "Context/correlation propagation, batching, retry rules, release gate, revision/ID checks, policy integration, and health behavior."),
    ("FastApiProviderContextTests", "The fastapi profile selects exactly one provider and makes readiness fail while liveness stays up during model outage."),
    ("ExcelReportServiceTests", "Original data, reusable colours, report columns, queue ordering, summary, legend, and reserved-sheet protection."),
    ("GlobalExceptionHandlerTests", "Private model/report failures map to safe public statuses and codes."),
]


PYTHON_TESTS = [
    ("tests/test_api.py", "Eleven FastAPI contract tests for one-time loading, live/ready separation, batch responses, validation, limits, correlation, runtime failures, and OpenAPI."),
    ("tests/test_evaluation.py", "Seven tests for dataset validity/reproducibility, release-gate failure, duplicate rejection, metrics, threshold selection, and candidate pinning."),
    ("tests/test_model_manifest.py", "Three tests for eligible/rejected candidates, commit pinning, complete label mapping, and native-label validation."),
]


def build_document():
    document = Document()
    configure_styles(document)
    bullet_num_id = add_numbering_definition(document, bullet=True)
    decimal_num_id = add_numbering_definition(document, bullet=False)
    props = document.core_properties
    props.title = "Content Filter Project Handbook"
    props.subject = "Beginner-friendly explanation of the Java and Python content-filter application"
    props.author = "Content Filter Project"
    props.keywords = "Spring Boot, FastAPI, Apache POI, Hugging Face, Excel, risk classification"

    # Cover: editorial_cover header pattern with compact-reference-guide body tokens.
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(108)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROJECT HANDBOOK")
    set_run_font(r, size=10.5, color=GOLD, bold=True)
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Content Filter")
    p = document.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("A beginner-friendly guide to every Java class, Python module, workflow, configuration, and test")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(42)
    r = p.add_run("Learning edition | Source state through Module 7 | 20 July 2026")
    set_run_font(r, size=10, color=MUTED, italic=True)

    add_table(document, ["Area", "Current state", "What that means"], [
        ["Java workflow", "Modules 1-3 and 6-7 implemented", "Upload, validation, policy, FastAPI adapter, and annotated workbook are working."],
        ["Python service", "Module 5 implemented", "Batch inference API and pinned provisional model runtime are implemented."],
        ["Model approval", "Module 4 gate open", "Synthetic smoke data is useful, but authorised human review/adjudication is still required."],
        ["Verification", "49 Java + 21 Python tests", "The implemented contracts currently pass automated verification."],
    ], [1750, 2200, 5410])
    add_callout(document, "Purpose of this handbook", "Read this when the class names begin to blur together. It explains the mental model first, then lets you look up every file without assuming prior Spring Boot, FastAPI, machine-learning, or Apache POI experience.", "info")
    document.add_page_break()

    document.add_heading("How to use this handbook", level=1)
    document.add_paragraph("You do not need to memorise every class. Start with Part I and the worked example. Then read the Java or Python part while keeping the architecture diagram visible. Use the class reference only when you open that file in the IDE.")
    for item in [
        "Part I: understand the product, two-service architecture, and complete request path.",
        "Part II: learn the Java/Spring Boot concepts and every Java class.",
        "Part III: learn the Python/FastAPI model service and evaluation pipeline.",
        "Part IV: understand the exact Java-Python JSON contract and configuration.",
        "Part V: run, test, debug, and safely continue development.",
        "Appendices: look up tests, errors, glossary terms, and a recommended code-reading order.",
    ]:
        add_number(document, item, decimal_num_id)
    add_callout(document, "One sentence mental model", "Java owns the user request, Excel, business decision, errors, and final report. Python owns model loading and score generation. The model does not make the final business decision.", "warning")

    document.add_heading("Contents", level=2)
    for item in [
        "Part I - What the application does and how one request moves",
        "Part II - Java layer and class-by-class reference",
        "Part III - Python layer, model runtime, and evaluation",
        "Part IV - Cross-service contract, configuration, and endpoints",
        "Part V - Running, testing, debugging, current limits, and next steps",
        "Appendix A - Test-class reference",
        "Appendix B - Glossary",
        "Appendix C - Recommended learning path",
    ]:
        add_bullet(document, item, bullet_num_id)

    document.add_page_break()
    document.add_heading("Part I - The big picture", level=1)
    document.add_heading("1. What problem is this project solving?", level=2)
    document.add_paragraph("The project is becoming a conversation-risk review platform. A user uploads an exported Excel workbook containing customer-support or AI-assisted messages. The system classifies each complete message, preserves the evidence and source coordinates, routes uncertain or serious cases to human review, and returns an annotated Excel workbook.")
    add_callout(document, "Important scope", "V1 classifies a complete sentence or message sequence. It does not highlight individual offensive words. NO_AUTOMATED_FLAG means no configured threshold was exceeded; it never means guaranteed safe.", "risk")
    document.add_heading("2. Why are Java and Python separate?", level=2)
    add_table(document, ["Java / Spring Boot", "Python / FastAPI"], [
        ["Public API and multipart upload", "Private internal model API"],
        ["Excel validation, extraction, coordinates, and report", "Hugging Face tokenizer/model loading and inference"],
        ["Thresholds, precedence, severity, and human-review decision", "Five provider-neutral risk scores and truncation evidence"],
        ["Problem Details, correlation, readiness, and download headers", "Pydantic validation, model readiness, and safe inference errors"],
        ["Stable business workflow even when model changes", "Replaceable candidate runtime and evaluation tooling"],
    ], [4680, 4680])
    document.add_paragraph("The separation is deliberate. Java is strong for enterprise APIs, contracts, and Apache POI. Python is the natural ecosystem for Hugging Face, PyTorch, and model evaluation. The RiskModel interface and JSON contract keep the boundary controlled.")

    document.add_heading("3. The complete versioned workflow", level=2)
    add_code(document, """User / Postman
    -> POST /api/v1/files/classify (multipart file)
    -> CorrelationIdFilter assigns X-Correlation-ID
    -> FileClassificationController
    -> WorkbookClassificationService
       -> ConversationRiskWorkflowService
          -> ExcelService extracts List<ExtractedSequence>
          -> RiskClassificationService
             -> RiskModel
                -> DeterministicRiskModel (default), OR
                -> FastApiRiskModel -> Python FastAPI -> Hugging Face model
             -> VersionedRiskPolicy creates ClassificationResult
          -> ConversationRiskAssessment keeps rows and decisions correlated
       -> ExcelReportService annotates original workbook
    -> ClassifiedWorkbook
    -> HTTP .xlsx download""")
    for item in [
        "Spring receives the file but the controller does not parse it.",
        "ExcelService validates the real file content, opens only supported OOXML, and creates one ExtractedSequence per mapped non-empty message.",
        "RiskClassificationService asks the selected RiskModel for scores and validates response IDs.",
        "VersionedRiskPolicy converts scores into a category, severity, confidence, and review decision.",
        "ExcelReportService uses saved coordinates to colour the source cell and add audit columns plus Review Queue, Summary, and Legend sheets.",
        "The controller returns bytes as an Excel attachment and forbids caching.",
    ]:
        add_number(document, item, decimal_num_id)

    document.add_heading("4. A worked example", level=2)
    add_table(document, ["conversation_id", "message_id", "speaker_role", "text", "language"], [
        ["C-17", "M-3", "customer", "You are stupid", "en"]
    ], [1600, 1400, 1450, 3460, 1450])
    add_labeled_paragraph(document, "Extraction", "ExcelService finds the text header, skips the header row, reads row index 1, and creates sequenceId sheet-0-row-1 with source column coordinates and context C-17/M-3/customer/en.")
    add_labeled_paragraph(document, "Model evidence", "The deterministic adapter would produce HARASSMENT_OR_INSULT=0.91 and low scores elsewhere. The FastAPI adapter would instead return the pinned Hugging Face model's five scores.")
    add_labeled_paragraph(document, "Policy decision", "The harassment threshold is 0.75, so 0.91 passes. VersionedRiskPolicy selects HARASSMENT_OR_INSULT, assigns HIGH severity, requires review, and records the reason and policy version.")
    add_labeled_paragraph(document, "Report", "ExcelReportService colours the original message cell orange, appends category/severity/confidence/review/model/policy/score columns, and places the message in Review Queue.")

    document.add_heading("5. The legacy path is different", level=2)
    add_code(document, "POST /upload -> UploadController -> UploadClassificationService -> ExcelService simple rows -> KeywordClassifier -> UploadResponse JSON")
    add_callout(document, "Do not confuse the two", "The legacy path returns abusive/professional/neutral JSON with random demonstration confidence. The versioned path returns an annotated workbook with the stable risk taxonomy and versioned policy.", "warning")

    document.add_page_break()
    document.add_heading("Part II - Java / Spring Boot layer", level=1)
    document.add_heading("6. Java concepts used in this project", level=2)
    concepts = [
        ("Controller", "The HTTP boundary. It reads request parts and creates HTTP responses, but should not contain Excel/model logic."),
        ("Service", "A class that performs or coordinates application work. Services can call other services through constructor-injected dependencies."),
        ("Domain record", "An immutable object representing a business concept such as ExtractedSequence or ClassificationResult."),
        ("DTO", "A transport shape used for HTTP JSON. The older DTOs belong to the legacy endpoint."),
        ("Interface / port", "A stable contract such as RiskModel. Different adapters can implement it without changing callers."),
        ("Adapter", "A concrete bridge to another implementation or system, such as FastApiRiskModel or DeterministicRiskModel."),
        ("Dependency injection", "Spring constructs classes and supplies their constructor dependencies. Code does not manually call new for application services."),
        ("Profile / conditional bean", "Configuration selects one implementation. The default is deterministic; the fastapi profile selects FastApiRiskModel."),
        ("Configuration properties", "Typed, validated settings bound from properties/environment variables."),
        ("Exception advice", "One global component maps internal exceptions to consistent safe HTTP errors."),
    ]
    add_table(document, ["Concept", "Meaning here"], [[a, b] for a, b in concepts], [2100, 7260])

    document.add_heading("7. Package map", level=2)
    add_code(document, """com.example.contentfilter
|-- ContentfilterApplication          starts Spring Boot
|-- config/                           typed deployment settings + FastAPI client
|-- controller/                       public HTTP endpoints
|-- domain/                           stable business records and enums
|-- dto/                              legacy JSON transport records
|-- exception/                        workbook/report/error translation
|-- service/                          workflows, policy, Excel, and model adapters
`-- web/                              request correlation filter""")

    for section_name, entries in JAVA_CLASSES.items():
        document.add_heading(section_name, level=2)
        for entry in entries:
            add_class_entry(document, bullet_num_id=bullet_num_id, **entry)

    document.add_page_break()
    document.add_heading("8. Deep dive: ExcelService", level=1)
    document.add_paragraph("ExcelService is long because file handling is security-sensitive. Its responsibilities can be understood as a pipeline:")
    for item in [
        "validateMetadata(): file exists, non-empty, within size, safe name, .xlsx extension, accepted MIME type.",
        "validateFileSignature(): checks the actual bytes begin as an OOXML ZIP package rather than trusting the name.",
        "new XSSFWorkbook(input): parses the workbook inside try-with-resources so streams are closed.",
        "validateWorkbook(): rejects macros, zero sheets, and too many sheets.",
        "resolveMappedColumns(): normalises header names, detects duplicate mapped headers, and finds required text plus optional context columns.",
        "extractFirstSheet(): walks physical rows, skips headers/blanks, formats cached cell values, creates coordinates/context, and enforces all limits.",
        "configurePoiSafeguards(): sets minimum inflate ratio and limits entry size, XML text, and part count to reduce zip-bomb risk.",
    ]:
        add_number(document, item, decimal_num_id)
    add_callout(document, "Why cached formulas?", "Recalculating formulas can change the workbook or invoke behavior the service does not need. V1 reads the cached displayed value and does not follow external links.", "info")

    document.add_heading("9. Deep dive: model evidence versus policy", level=2)
    add_code(document, """ModelPrediction scores
    -> input truncated?              -> MANUAL_REVIEW / HIGH
    -> 2 passing scores too close?   -> MANUAL_REVIEW / highest severity
    -> any score passes threshold?   -> first category in precedence
    -> score near a threshold?       -> MANUAL_REVIEW / LOW
    -> otherwise                     -> NO_AUTOMATED_FLAG / NONE""")
    add_callout(document, "Key design rule", "FastAPI never returns severity or reviewRequired. Those are business decisions controlled by versioned Java policy.", "warning")

    document.add_heading("10. Deep dive: annotated workbook", level=2)
    add_table(document, ["Primary category", "Typical severity", "Excel fill"], [
        ["THREAT", "CRITICAL", "Red"],
        ["HATE_OR_IDENTITY_ATTACK", "HIGH", "Coral"],
        ["HARASSMENT_OR_INSULT", "HIGH", "Orange"],
        ["OBSCENE_OR_PROFANE", "MEDIUM", "Yellow"],
        ["GENERAL_TOXICITY", "MEDIUM", "Light orange"],
        ["MANUAL_REVIEW", "Policy-assigned", "Light blue"],
        ["NO_AUTOMATED_FLAG", "NONE", "Light green"],
    ], [3600, 2400, 3360])
    document.add_paragraph("The source worksheet receives explicit text columns for primary category, severity, confidence, review status/reason, model ID/revision, policy version, and all five category scores. Colour is never the only explanation. The additional sheets are:")
    for item in [
        "Review Queue: only review-required messages, sorted by severity and then confidence.",
        "Summary: counts by primary category and review status.",
        "Legend: category colour meanings, typical severity, and explicit limitations.",
    ]:
        add_bullet(document, item, bullet_num_id)

    document.add_page_break()
    document.add_heading("Part III - Python / FastAPI layer", level=1)
    document.add_heading("11. Python service mental model", level=2)
    add_code(document, """FastAPI request JSON
    -> Pydantic schemas validate shape, sizes, uniqueness, English V1
    -> ClassificationApplicationService enforces runtime limits
    -> HuggingFaceRiskRuntime tokenizes and runs pinned model
    -> native model labels map to five RiskCategory scores
    -> application service checks IDs and scores
    -> FastAPI response returns evidence + exact model revision
    -> Java validates it and applies business policy""")
    add_callout(document, "Why runtime injection exists", "create_app() accepts a runtime_factory so API tests can use a tiny fake runtime. Tests therefore do not download model weights or require PyTorch inference.", "info")

    document.add_heading("12. FastAPI module-by-module reference", level=2)
    for path, explanation in PYTHON_MODULES:
        document.add_heading(path, level=3)
        p = document.add_paragraph(style="File Path")
        p.add_run("model-service/" + path)
        document.add_paragraph(explanation)

    document.add_heading("13. The Python API endpoints", level=2)
    add_table(document, ["Endpoint", "Meaning", "Does not mean"], [
        ["GET /live", "The FastAPI process is running.", "The model is loaded."],
        ["GET /ready", "The pinned model/tokenizer loaded and can infer.", "The model is production approved."],
        ["POST /api/v1/classify/batch", "Returns correlated five-score evidence.", "A final risk/severity/review decision."],
        ["GET /docs", "Interactive OpenAPI documentation.", "A public internet API."],
    ], [2500, 3430, 3430])

    document.add_heading("14. What happens during Hugging Face inference?", level=2)
    for item in [
        "The manifest selects one eligible model by key and immutable commit hash.",
        "Tokenizer and model are downloaded/loaded with trust_remote_code=False; the model requires safetensors.",
        "The runtime verifies that configured native labels actually exist in model.config.id2label.",
        "Messages are tokenised in small inference batches, padded, and truncated to at most 512 tokens.",
        "A separate pre-tokenisation check records whether the original input exceeded that token limit.",
        "torch.inference_mode() disables gradient tracking because the service predicts rather than trains.",
        "Sigmoid converts independent multi-label logits to probabilities.",
        "Native probabilities map to each domain category; where several native labels map to one category, the maximum is used.",
        "A lock prevents concurrent access to the current CPU runtime.",
    ]:
        add_number(document, item, decimal_num_id)

    document.add_heading("15. Evaluation scripts", level=2)
    for path, explanation in EVALUATION_SCRIPTS:
        document.add_heading(path, level=3)
        p = document.add_paragraph(style="File Path")
        p.add_run("model-service/" + path)
        document.add_paragraph(explanation)

    document.add_heading("16. Evaluation data and artifacts", level=2)
    add_table(document, ["Artifact", "Purpose", "Current truth status"], [
        ["v1-seed.jsonl", "Original 40-row tooling seed.", "Synthetic draft"],
        ["v1-domain-synthetic.jsonl", "Reproducible 240-row realistic smoke/regression corpus.", "Synthetic draft"],
        ["v1-domain-adjudicated.jsonl", "Local candidate file named as adjudicated.", "Untracked and not verified; do not treat as human ground truth"],
        ["candidate-models.json", "Pinned candidate identity, eligibility, licence, and label mapping.", "Versioned candidate manifest"],
        ["baseline-thresholds.json", "Exercises metric tooling at 0.5.", "Not production approved"],
        ["ANNOTATION_GUIDE.md", "Defines human multi-label categories, context rules, and review process.", "Process guidance"],
        ["SMOKE_REPORT.md / DOMAIN_SMOKE_REPORT.md", "Record proof runs and observed failure modes.", "No production model selected"],
    ], [2700, 3850, 2810])
    add_callout(document, "Release gate", "Two people must label authorised, privacy-reviewed data independently; an authorised third reviewer adjudicates disagreements. Only a frozen adjudicated dataset may support a production model decision.", "risk")

    document.add_page_break()
    document.add_heading("Part IV - Java/Python contract and configuration", level=1)
    document.add_heading("17. The internal batch request", level=2)
    add_code(document, """POST http://127.0.0.1:8000/api/v1/classify/batch
{
  "sequences": [{
    "sequenceId": "sheet-0-row-1",
    "text": "You are stupid",
    "conversationId": "C-17",
    "speakerRole": "customer",
    "language": "en"
  }]
}""")
    document.add_heading("18. The internal batch response", level=2)
    add_code(document, """{
  "model": "minuva/MiniLMv2-toxic-jigsaw",
  "revision": "00eacca7ba7c09b1e82db508b03a901bf9cc89eb",
  "evaluationStatus": "PROVISIONAL",
  "approvedForProduction": false,
  "predictions": [{
    "sequenceId": "sheet-0-row-1",
    "scores": {
      "THREAT": 0.01,
      "HATE_OR_IDENTITY_ATTACK": 0.02,
      "HARASSMENT_OR_INSULT": 0.91,
      "OBSCENE_OR_PROFANE": 0.03,
      "GENERAL_TOXICITY": 0.66
    },
    "inputTruncated": false
  }]
}""")
    document.add_paragraph("FastApiRiskModel validates the model/revision/release metadata and prediction IDs before creating Java ModelPrediction objects. RiskClassificationService validates correlation again at the application layer. This intentional duplication protects separate trust boundaries.")

    document.add_heading("19. Configuration files", level=2)
    add_table(document, ["File", "Purpose"], [
        ["pom.xml", "Spring Boot parent, Java 21, Web MVC, Actuator, Validation, Apache POI 5.5.1, test support, and executable JAR packaging."],
        ["application.properties", "Default deterministic provider, upload limits, policy thresholds/precedence, FastAPI identity/timeouts, graceful shutdown, and minimal Actuator exposure."],
        ["application-fastapi.properties", "Development-only profile selecting FastAPI, allowing provisional evidence, and adding model health to readiness."],
        ["model-service/requirements-api.txt", "Pinned FastAPI/Pydantic/Starlette/Uvicorn runtime."],
        ["model-service/requirements-model.txt", "Pinned PyTorch and Transformers model runtime."],
        ["model-service/requirements-test.txt", "Lightweight API/test dependencies without model weights."],
        [".github/workflows/ci.yml", "Runs reproducible dataset checks, 21 Python tests/source compilation, and Maven Java verification on pushes/PRs."],
    ], [3000, 6360])

    document.add_heading("20. Important property groups", level=2)
    add_table(document, ["Prefix", "Controls"], [
        ["spring.servlet.multipart", "Web-layer request/file maximums before Apache POI runs."],
        ["content-filter.upload", "Workbook, row, sequence, cell, text, and time limits."],
        ["content-filter.classification.policy", "Policy version, uncertainty margin, category thresholds, and precedence."],
        ["content-filter.classification.model-service", "FastAPI origin, expected model/revision, timeouts, batch size, attempts, and provisional waiver."],
        ["management.*", "Actuator endpoints and readiness composition."],
        ["CONTENT_FILTER_MODEL_*", "Python model candidate, cache, batch, and text limits."],
    ], [3600, 5760])

    document.add_heading("21. Correlation and error contract", level=2)
    add_code(document, """Client X-Correlation-ID
    -> Java CorrelationIdFilter validates/generates ID
    -> Java logs and Problem Details use same ID
    -> FastApiRiskModel forwards ID
    -> Python middleware preserves/generates ID
    -> Python response returns ID""")
    add_table(document, ["Example error", "HTTP", "Public errorCode"], [
        ["Missing/empty multipart file", "400", "INVALID_UPLOAD"],
        ["Wrong extension, MIME, or non-OOXML signature", "415", "UNSUPPORTED_WORKBOOK_TYPE"],
        ["Workbook/resource limit exceeded", "413", "WORKBOOK_LIMIT_EXCEEDED or UPLOAD_TOO_LARGE"],
        ["Reserved report sheet / malformed workbook", "400", "INVALID_WORKBOOK"],
        ["FastAPI unavailable", "503", "MODEL_SERVICE_UNAVAILABLE"],
        ["FastAPI timeout", "504", "MODEL_SERVICE_TIMEOUT"],
        ["Wrong model/revision or invalid evidence", "502", "INVALID_MODEL_SERVICE_RESPONSE"],
        ["Provisional model without waiver", "503", "MODEL_NOT_APPROVED"],
        ["Report-writing failure", "500", "REPORT_GENERATION_FAILED"],
    ], [4500, 950, 3910])

    document.add_page_break()
    document.add_heading("Part V - Run, test, and debug", level=1)
    document.add_heading("22. Run Java with the deterministic development adapter", level=2)
    add_code(document, ".\\mvnw.cmd spring-boot:run")
    document.add_paragraph("This starts only Java on port 8080. The versioned endpoint works without Python because DeterministicRiskModel is the default provider.")

    document.add_heading("23. Run the actual FastAPI bridge", level=2)
    add_code(document, """# From repository root
python -m venv .venv
& .\\.venv\\Scripts\\python.exe -m pip install -r model-service/requirements.txt
& .\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --app-dir model-service --host 127.0.0.1 --port 8000

# In another terminal after GET /ready returns READY
.\\mvnw.cmd spring-boot:run "-Dspring-boot.run.profiles=fastapi""")
    add_callout(document, "Development waiver", "application-fastapi.properties permits approvedForProduction=false only for local development while Module 4 remains open. Production configuration must not enable this waiver.", "warning")

    document.add_heading("24. Test in Postman", level=2)
    for item in [
        "Create POST http://localhost:8080/api/v1/files/classify.",
        "Choose Body -> form-data.",
        "Add key file, change its type from Text to File, and select an .xlsx workbook.",
        "The first worksheet must have a first-row header named text. Optional standard headers are conversation_id, message_id, speaker_role, timestamp, language, and channel.",
        "Send. On HTTP 200, use Save Response to file and open classified-<original>.xlsx.",
        "Inspect the source row, Content Filter columns, Review Queue, Summary, and Legend.",
        "Use response X-Correlation-ID when troubleshooting an error.",
    ]:
        add_number(document, item, decimal_num_id)

    document.add_heading("25. Run all automated tests", level=2)
    add_code(document, """.\\mvnw.cmd clean package
& .\\.venv\\Scripts\\python.exe -m pip check
& .\\.venv\\Scripts\\python.exe -m pytest model-service/tests -q""")
    add_labeled_paragraph(document, "Verified result at this snapshot", "49 Java tests and 21 Python tests pass; the executable Spring Boot JAR is produced.")

    document.add_heading("26. Debugging map", level=2)
    add_table(document, ["Symptom", "Start here", "Likely area"], [
        ["400 INVALID_UPLOAD", "Request file part and filename", "Controller boundary / ExcelService metadata"],
        ["400 INVALID_WORKBOOK", "Headers, reserved sheets, workbook readability", "ExcelService or ExcelReportService"],
        ["413", "Configured file/row/text limits", "UploadLimitsProperties / application.properties"],
        ["FastAPI /live works but /ready fails", "Model dependencies, cache, manifest, pinned revision", "Python lifespan / HuggingFaceRiskRuntime"],
        ["Spring liveness UP but readiness DOWN", "FastAPI /ready and Java fastapi profile", "FastApiModelHealthIndicator"],
        ["502 invalid model response", "Model ID/revision, prediction IDs, categories, release metadata", "FastApiRiskModel contract validation"],
        ["Unexpected category", "Raw scores, thresholds, precedence, uncertainty margin", "VersionedRiskPolicy"],
        ["Wrong Excel cell coloured", "sequenceId and coordinates", "ExcelService -> ConversationRiskAssessment -> ExcelReportService"],
        ["Missing row in Review Queue", "manualReviewRequired and result correlation", "VersionedRiskPolicy / ExcelReportService"],
    ], [2800, 3100, 3460])

    document.add_heading("27. What is completed and what is not", level=2)
    add_table(document, ["Area", "Status"], [
        ["Modules 1-3", "Operational baseline, safe extraction, mapping/domain/policy implemented."],
        ["Module 4", "Tooling and synthetic corpus implemented; authorised human review/adjudication and production model selection still open."],
        ["Modules 5-7", "FastAPI inference, Spring adapter, and annotated workbook endpoint implemented and tested."],
        ["Module 8", "Broader security, load, contract, and live cross-service integration hardening is next."],
        ["Modules 9-10", "Docker/observability/hardening, upload UI, and controlled V1 release remain."],
        ["Future V2-V5", "Authentication/analytics, enterprise webhooks, multilingual support, and user-configurable rules are not implemented yet."],
    ], [2200, 7160])

    document.add_page_break()
    document.add_heading("Appendix A - Automated test reference", level=1)
    document.add_heading("Java test classes", level=2)
    add_table(document, ["Test class", "What it protects"], [[a, b] for a, b in JAVA_TESTS], [3300, 6060])
    document.add_heading("Python test modules", level=2)
    add_table(document, ["Test module", "What it protects"], [[a, b] for a, b in PYTHON_TESTS], [3300, 6060])
    add_callout(document, "What tests cannot prove", "Passing tests prove the implemented contracts behave as expected. They do not prove that a provisional model is fair, accurate, or safe on real production conversations. That requires the Module 4 human evaluation gate.", "risk")

    document.add_heading("Appendix B - Glossary", level=1)
    glossary = [
        ("Actuator", "Spring Boot operational endpoints for health, liveness, readiness, and info."),
        ("Adapter", "A concrete implementation that connects an application interface to a provider or technology."),
        ("Adjudication", "An authorised third review that resolves disagreements between independent labelers."),
        ("Apache POI", "Java library used to read and write Microsoft Office files, including .xlsx."),
        ("Batch", "Several messages sent to the model in one request for efficiency and a stable contract."),
        ("Bean", "An object created and managed by the Spring container."),
        ("Calibration split", "Evaluation examples used to choose thresholds. It must not be mixed with the final test split."),
        ("Confidence", "The policy's selected score or inverse risk evidence; it is not automatically a calibrated real-world probability."),
        ("Correlation ID", "A safe request identifier copied across responses and services to join logs and errors."),
        ("DTO", "Data Transfer Object used primarily to shape data crossing an API boundary."),
        ("Domain object", "A record or enum representing an application concept independent of HTTP or a model provider."),
        ("FastAPI", "Python web framework exposing the internal model inference API."),
        ("Hugging Face", "Model ecosystem used to obtain the pinned tokenizer and sequence-classification model."),
        ("Immutable revision", "A commit hash identifying exact model artifacts rather than a moving branch name."),
        ("Inference", "Running a trained model to obtain predictions; different from training."),
        ("Interface / port", "A Java contract that callers depend on while implementations can be swapped."),
        ("JSONL", "JSON Lines: one JSON object per line, used for datasets and predictions."),
        ("Liveness", "Whether the process is running. It does not promise dependencies are ready."),
        ("Logit", "Raw model output converted to multi-label probability-like scores with sigmoid."),
        ("MDC", "Mapped Diagnostic Context: per-request logging data, used here for correlationId."),
        ("MultipartFile", "Spring's representation of an uploaded file part in multipart/form-data."),
        ("Multi-label", "A message may support several risk categories at the same time."),
        ("OOXML", "The ZIP/XML format behind modern .xlsx files."),
        ("Pydantic", "Python validation library used for strict API schemas and settings."),
        ("Policy", "Deterministic rules converting model evidence into business outcomes."),
        ("Profile", "A named Spring configuration overlay such as fastapi."),
        ("Provider-neutral", "Application concepts that do not expose one model's native label names."),
        ("Readiness", "Whether the service and selected dependency are able to accept work safely."),
        ("Record", "Compact Java immutable data carrier with generated accessors and value equality."),
        ("RestClient", "Spring HTTP client used by Java to call FastAPI."),
        ("Safetensors", "Model-weight format required here to reduce unsafe deserialisation risk."),
        ("Sequence", "The V1 classification unit: one complete mapped message or fallback non-empty row."),
        ("Sigmoid", "Function converting each independent multi-label logit into a value from 0 to 1."),
        ("Spring dependency injection", "Spring creates application classes and passes constructor dependencies automatically."),
        ("Test split", "Held-out examples used once after threshold selection to estimate candidate performance."),
        ("Threshold", "Configured score boundary above which a category is treated as detected."),
        ("Truncation", "Model token limit removed part of the message; policy requires human review of the complete text."),
        ("Zip bomb", "A tiny compressed file that expands massively; Apache POI safeguards reduce this resource attack."),
    ]
    add_table(document, ["Term", "Meaning in this project"], [[a, b] for a, b in glossary], [2400, 6960])

    document.add_heading("Appendix C - Recommended learning path", level=1)
    learning_steps = [
        "Open FileClassificationController. Notice that it has one dependency and no Excel/model logic.",
        "Follow WorkbookClassificationService.classify() to see the whole use case in a few lines.",
        "Read ConversationRiskWorkflowService, then examine ExtractedSequence and ConversationRiskAssessment.",
        "Read ExcelService.extractConversationSequences() and trace one worksheet row into an ExtractedSequence.",
        "Read RiskModel, then compare DeterministicRiskModel and FastApiRiskModel.",
        "Read RiskClassificationService to understand correlation by sequence ID.",
        "Read VersionedRiskPolicy with VersionedRiskPolicyTests side by side.",
        "Read ExcelReportService.generate() and ExcelReportServiceTests together.",
        "Move to Python schemas.py, service.py, runtime.py, then main.py in that order.",
        "Read test_api.py to see how a fake runtime proves the HTTP contract without model downloads.",
        "Finally read evaluation_lib.py and the annotation guide to understand why model approval is separate from code completion.",
    ]
    for step in learning_steps:
        add_number(document, step, decimal_num_id)
    add_callout(document, "Learning strategy", "For each class ask four questions: What data enters? What data leaves? Which dependency does the work? Which test proves the behavior? If you can answer those, you understand the class even before every private method is familiar.", "info")

    document.add_heading("Final architecture recap", level=2)
    add_code(document, """PUBLIC SIDE (Java)                         PRIVATE SIDE (Python)
Multipart Excel
  -> validate + extract coordinates
  -> batch RiskModel -----------------------> strict FastAPI batch schema
                                               -> pinned model runtime
                                               -> five risk scores
  <- validate identity + scores <-----------
  -> versioned Java policy
  -> annotated Excel report
  -> non-cacheable download""")
    document.add_paragraph("The architecture is intentionally layered so each concern has one home: controllers own HTTP, Excel services own workbook processing, domain records protect data meaning, interfaces protect boundaries, Python owns inference, Java owns policy, and tests protect each handoff.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
